#!/usr/bin/env python3
"""exam-prep ingest: read every source file once, cheaply.

For each input it extracts the text layer page by page and decides which pages actually need to be
*looked at* (scanned pages, garbled formulas, figures, photos). Only those pages are rendered — at low
resolution and packed several to an image — so the model spends image tokens where text extraction fails,
not on every page.

Usage:
    python ingest.py slides.pdf ch3.pptx notes.jpg exam2025.pdf --out ingest/
    python ingest.py "textbook.pdf:120-180,201-230" --out ingest/     only these PDF pages (PDF numbering, not printed)

A PDF longer than --book-pages (a whole textbook) is not extracted when no pages are given: the report points to its
bookmarks with page ranges instead, so the scope can be chosen first and only those chapters read with "file.pdf:a-b".

Writes into --out:
    report.md              what was found, which pages to view, and what to read next   <- read this first
    text/<file>.txt        extracted text with "=== <file> p.N ===" page markers
    text/<file>-outline.md table of contents of a whole book (instead of its text)
    sheets/<file>-K.png    contact sheets of the pages that need viewing (2x2 slides or 2x1 portrait pages)

Supported: .pdf .pptx .docx .md .txt .png .jpg .jpeg .webp .heic(if Pillow can open it).
Needs: poppler (pdftotext, pdftoppm) for PDFs; python-pptx for .pptx; Pillow for images and sheets.
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    sys.exit("Pillow is required: pip install --break-system-packages pillow")

SHEET_MAX = 1568  # long side of a contact sheet; larger images get downscaled by the model anyway
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".heic", ".heif"}
MATH_ALNUM = re.compile(r"[\U0001D400-\U0001D7FF]")  # 𝐴 𝜆 … : math italic letters from PDF formulas
MATH_SYM = re.compile(r"[∑∏∫√∞≤≥≠≈±∓×÷∂∇λμσθπαβγδεΣΠΔ⁻¹²³ᵀ→⇒⇔∈∉⊂⊆∪∩]")
# bold fonts for page labels: Linux, Windows, macOS
LABEL_FONTS = ["/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
               "C:/Windows/Fonts/arialbd.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"]
# LibreOffice is usually not on PATH on Windows and macOS
SOFFICE_PATHS = ["C:/Program Files/LibreOffice/program/soffice.exe", "C:/Program Files (x86)/LibreOffice/program/soffice.exe",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"]


@dataclass
class Page:
    n: int
    text: str
    reasons: list[str] = field(default_factory=list)


@dataclass
class Source:
    path: Path
    kind: str
    pages: list[Page] = field(default_factory=list)
    note: str = ""
    render_pdf: Path | None = None  # PDF used to render pages that need viewing
    wanted: list[int] | None = None  # pages asked for with "file.pdf:a-b"; None = all
    spec: str = ""                  # those pages as written, for the report
    outline: str = ""               # set instead of pages for a whole book
    pages_total: int = 0            # page count of that book
    image: Path | None = None       # for photos


# ----------------------------------------------------------------------------- heuristics
def visual_reasons(text: str, has_images: bool = False) -> list[str]:
    """Why a page's text layer is not enough on its own."""
    t = text.strip()
    reasons = []
    chars = len(re.sub(r"\s", "", t))
    if chars < 40:
        reasons.append("幾乎沒有文字（掃描頁或純圖）")
    lines = [l for l in t.splitlines() if l.strip()]
    if lines:
        frag = sum(1 for l in lines if len(l.strip()) <= 3) / len(lines)
        if frag > 0.3 and len(lines) >= 6:
            reasons.append("文字破碎（多半是公式或矩陣被拆開）")
    if len(MATH_ALNUM.findall(t)) >= 3:
        reasons.append("含 PDF 公式字元（公式結構可能遺失）")
    elif len(MATH_SYM.findall(t)) >= 8 and chars < 600:
        reasons.append("公式密集")
    if has_images and chars < 300:
        reasons.append("有圖片且文字少（圖表或示意圖）")
    return reasons


def run(cmd: list[str]) -> str:
    # poppler and pandoc write UTF-8; without an explicit encoding Windows would decode with the locale (e.g. cp950)
    return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", check=True).stdout


def poppler(tool: str) -> str:
    """Path to a poppler tool. Git for Windows puts xpdf's pdftotext first on PATH; xpdf drops CJK text and
    mangles math symbols, so take the copy next to pdfinfo (which Git does not ship) when there is one."""
    info = shutil.which("pdfinfo")
    path = (shutil.which(tool, path=str(Path(info).parent)) if info else None) or shutil.which(tool)
    if not path:
        raise RuntimeError(f"找不到 {tool}，請安裝 poppler（執行 setup.py --check 看安裝方式）")
    return path


def find_soffice() -> str | None:
    return shutil.which("soffice") or next((p for p in SOFFICE_PATHS if Path(p).exists()), None)


# ----------------------------------------------------------------------------- readers
def pdf_pages_with_images(pdf: Path) -> set[int]:
    try:
        out = run([poppler("pdfimages"), "-list", str(pdf)])
    except (RuntimeError, subprocess.CalledProcessError):
        return set()
    pages = set()
    for line in out.splitlines()[2:]:
        parts = line.split()
        if parts and parts[0].isdigit():
            w, h = (int(parts[3]), int(parts[4])) if len(parts) > 4 and parts[3].isdigit() else (0, 0)
            if w * h >= 150 * 150:  # ignore logos / bullets
                pages.add(int(parts[0]))
    return pages


def pdf_outline(pdf: Path, n: int) -> str:
    """Bookmarks as a Markdown table with each heading's page range; without bookmarks, the text of the first
    pages, where the printed table of contents usually is."""
    try:
        xml = run([poppler("pdftohtml"), "-xml", "-i", "-q", "-f", "1", "-l", "1", "-stdout", str(pdf)])
        root = ET.fromstring(re.sub(r"<!DOCTYPE[^>]*>", "", xml))
    except (RuntimeError, subprocess.CalledProcessError, ET.ParseError):
        root = None
    items: list[tuple[int, int, str]] = []  # (level, page, title)

    def walk(node, level):
        for child in node:
            if child.tag == "item" and child.get("page", "").isdigit():
                items.append((level, int(child.get("page")), " ".join((child.text or "").split())))
            elif child.tag == "outline":
                walk(child, level + 1)

    outline = root.find("outline") if root is not None else None
    if outline is not None:
        walk(outline, 1)
    if not items:
        first = min(n, 30)
        text = run([poppler("pdftotext"), "-enc", "UTF-8", "-layout", "-f", "1", "-l", str(first), str(pdf), "-"])
        return (f"這份 PDF 沒有書籤，以下是前 {first} 頁的文字，目錄通常在這裡。目錄上印的是書本頁碼，和 PDF 頁碼"
                f"通常差一個固定值：抽一頁（`pdftotext -f 40 -l 40 檔案.pdf -`）對照章名就能算出來。\n\n{text}")
    lines = ["| 層級 | 標題 | PDF 頁碼 |", "|---|---|---|"]
    for k, (level, page, title) in enumerate(items):
        if level > 2:
            continue
        # a heading runs until the next heading at the same or a higher level
        end = next((pg for lv, pg, _ in items[k + 1:] if lv <= level), n + 1) - 1
        span = f"{page}-{end}" if end > page else str(page)
        lines.append(f"| {level} | {'　' * (level - 1)}{title} | {span} |")
    return "\n".join(lines)


def read_pdf(src: Source, book_pages: int, name: str):
    info = run([poppler("pdfinfo"), str(src.path)])
    n = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    if src.wanted is None and n > book_pages:
        src.outline, src.pages_total = pdf_outline(src.path, n), n
        src.note = (f"整本書（{n} 頁），沒有抽全文。先用 `text/{name}-outline.md` 的目錄問使用者範圍，"
                    f"再用 `\"{src.path}:起-迄\"`（PDF 頁碼，可用逗號接多段）只讀選定的頁")
        return
    if src.wanted and src.wanted[-1] > n:
        src.note = f"這份 PDF 只有 {n} 頁，超出的頁碼已略過"
    imgs = pdf_pages_with_images(src.path)
    pdftotext = poppler("pdftotext")
    for i in (pg for pg in (src.wanted or range(1, n + 1)) if 1 <= pg <= n):
        # -enc UTF-8 is poppler's default but not xpdf's (Latin-1), in case only xpdf is installed
        text = run([pdftotext, "-enc", "UTF-8", "-layout", "-f", str(i), "-l", str(i), str(src.path), "-"])
        src.pages.append(Page(i, text, visual_reasons(text, i in imgs)))
    src.render_pdf = src.path


def read_pptx(src: Source, work: Path):
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "--break-system-packages", "-q", "python-pptx"])
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(src.path))
    for i, slide in enumerate(prs.slides, 1):
        texts, has_visual = [], False
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text)
            if sh.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.CHART):
                has_visual = True
            if getattr(sh, "has_table", False) and sh.has_table:
                texts.append("\n".join(" | ".join(c.text for c in r.cells) for r in sh.table.rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("[講者備註] " + slide.notes_slide.notes_text_frame.text)
        text = "\n".join(texts)
        reasons = visual_reasons(text, has_visual)
        if has_visual and not any("圖" in r for r in reasons):
            reasons.append("含圖片、圖表或嵌入公式物件")
        src.pages.append(Page(i, text, reasons))
    # render via LibreOffice only if some slide needs viewing
    soffice = find_soffice()
    if any(p.reasons for p in src.pages) and soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(src.path)],
                       capture_output=True)
        pdf = work / (src.path.stem + ".pdf")
        if pdf.exists():
            src.render_pdf = pdf
        else:
            src.note = "LibreOffice 轉檔失敗，需要看圖的投影片請另外處理"
    elif any(p.reasons for p in src.pages):
        src.note = "沒有 LibreOffice（soffice），無法把需要看圖的投影片轉成圖片；可請使用者另存成 PDF 上傳"


def read_docx(src: Source):
    if shutil.which("pandoc"):
        text = run(["pandoc", str(src.path), "-t", "gfm", "--wrap=none"])
    else:
        try:
            import docx  # python-docx
        except ImportError:
            subprocess.run([sys.executable, "-m", "pip", "install", "--break-system-packages", "-q", "python-docx"])
            import docx
        text = "\n".join(p.text for p in docx.Document(str(src.path)).paragraphs)
    src.pages.append(Page(1, text, []))
    import zipfile
    try:
        xml = zipfile.ZipFile(src.path).read("word/document.xml").decode("utf-8", "ignore")
    except Exception:  # noqa: BLE001
        xml = ""
    if "<m:oMath" in xml and "$" not in text:
        src.note = "文件含 Word 方程式，轉出的文字可能缺公式；必要時請使用者另存 PDF"


def read_image(src: Source, work: Path):
    im = Image.open(src.path)
    im = im.convert("RGB")
    scale = min(1.0, SHEET_MAX / max(im.size))
    if scale < 1:
        im = im.resize((int(im.width * scale), int(im.height * scale)), Image.LANCZOS)
    out = work / f"{src.path.stem}.jpg"
    im.save(out, quality=85)
    src.image = out
    src.pages.append(Page(1, "", ["照片／圖片，需要直接看"]))


# ----------------------------------------------------------------------------- sheets
def render_pages(pdf: Path, pages: list[int], work: Path, dpi: int) -> list[tuple[int, Image.Image]]:
    out = []
    for n in pages:
        prefix = work / f"r{n}"
        subprocess.run([poppler("pdftoppm"), "-r", str(dpi), "-f", str(n), "-l", str(n), "-png", "-singlefile", str(pdf), str(prefix)],
                       check=True, capture_output=True)
        out.append((n, Image.open(f"{prefix}.png").convert("RGB")))
    return out


def label_font(size=22):
    for p in LABEL_FONTS:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    try:
        return ImageFont.load_default(size)  # Pillow >= 10.1
    except TypeError:
        return ImageFont.load_default()


def make_sheets(name: str, imgs: list[tuple[int, Image.Image]], outdir: Path) -> list[Path]:
    if not imgs:
        return []
    landscape = imgs[0][1].width >= imgs[0][1].height
    cols, rows = (2, 2) if landscape else (2, 1)
    per = cols * rows
    cell_w = (SHEET_MAX - 10 * (cols + 1)) // cols
    font = label_font()
    paths = []
    for k in range(0, len(imgs), per):
        chunk = imgs[k:k + per]
        cells = []
        for n, im in chunk:
            s = cell_w / im.width
            cells.append((n, im.resize((cell_w, int(im.height * s)), Image.LANCZOS)))
        cell_h = max(c.height for _, c in cells)
        used = min(cols, len(cells))  # don't pad a lone page with an empty column
        r = (len(cells) + cols - 1) // cols
        sheet = Image.new("RGB", (used * cell_w + 10 * (used + 1), r * (cell_h + 34) + 10), "#9aa0a6")
        d = ImageDraw.Draw(sheet)
        for idx, (n, c) in enumerate(cells):
            x = 10 + (idx % cols) * (cell_w + 10)
            y = 10 + (idx // cols) * (cell_h + 34)
            d.text((x, y), f"p.{n}", fill="white", font=font)
            sheet.paste(c, (x, y + 28))
        p = outdir / f"{name}-{k // per + 1}.png"
        sheet.save(p, optimize=True)
        paths.append(p)
    return paths


# ----------------------------------------------------------------------------- main
def safe(name: str) -> str:
    return re.sub(r"[^\w一-鿿.-]+", "_", name)


def parse_target(arg: str) -> tuple[Path, list[int] | None, str]:
    """'book.pdf:120-180,200' -> (path, pages, '120-180,200'). An existing path wins, so 'D:/x.pdf' stays a path."""
    m = re.fullmatch(r"(.+):(\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*)", arg)
    if Path(arg).exists() or not m:
        return Path(arg), None, ""
    pages: list[int] = []
    for part in m.group(2).split(","):
        lo, _, hi = part.partition("-")
        pages += range(int(lo), int(hi or lo) + 1)
    return Path(m.group(1)), sorted(set(pages)), m.group(2)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", help='files; "file.pdf:120-180,200" reads only those PDF pages')
    ap.add_argument("--out", type=Path, default=Path("ingest"))
    ap.add_argument("--dpi", type=int, default=72, help="render resolution for pages that need viewing")
    ap.add_argument("--max-visual", type=int, default=40, help="cap on rendered pages per file")
    ap.add_argument("--book-pages", type=int, default=150,
                    help="a PDF longer than this, given without pages, is treated as a whole book (outline only)")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # the report has symbols a cp950 console cannot print

    (a.out / "text").mkdir(parents=True, exist_ok=True)
    (a.out / "sheets").mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="ingest-"))
    sources: list[Source] = []

    names: list[str] = []
    for arg in a.files:
        f, wanted, spec = parse_target(arg)
        if not f.exists():
            print(f"skip (not found): {f}", file=sys.stderr)
            continue
        ext = f.suffix.lower()
        src = Source(f, ext.lstrip("."), wanted=wanted, spec=spec)
        # the range goes into the name: the same book may be given twice with different chapters
        name = safe(f.stem) + (f"_p{spec.replace(',', '_')}" if spec else "")
        try:
            if ext == ".pdf":
                read_pdf(src, a.book_pages, name)
            elif ext == ".pptx":
                read_pptx(src, work)
            elif ext == ".docx":
                read_docx(src)
            elif ext in IMG_EXT:
                read_image(src, work)
            elif ext in {".md", ".txt", ".markdown"}:
                src.pages.append(Page(1, f.read_text(encoding="utf-8", errors="replace"), []))
            else:
                src.note = f"不支援的格式 {ext}，請手動處理"
        except Exception as e:  # noqa: BLE001
            src.note = f"讀取失敗：{e}"
        if wanted is not None and ext != ".pdf":
            src.pages = [pg for pg in src.pages if pg.n in wanted]
        sources.append(src)
        names.append(name)

    lines = ["# Ingest report", "",
             "先讀這份報告，再依序：① 讀 `text/` 的文字 ② **只看**下表列出的 sheets ③ 寫 `materials.md`。",
             "沒列在 sheets 的頁面，文字層已足夠，不需要再看圖。", "",
             "| 檔案 | 類型 | 頁數 | 文字字數 | 需要看圖的頁 | sheets |", "|---|---|---|---|---|---|"]
    total_sheets = 0
    for src, name in zip(sources, names):
        if src.outline:
            (a.out / "text" / f"{name}-outline.md").write_text(
                f"# {src.path.name} 目錄（共 {src.pages_total} 頁）\n\n{src.outline}\n", encoding="utf-8")
            lines.append(f"| {src.path.name} | 整本書 | {src.pages_total} | — | — | — |")
            lines.append(f"| ↳ 注意 | {src.note} | | | | |")
            continue
        if src.pages:
            body = "\n\n".join(f"=== {src.path.name} p.{p.n} ===\n{p.text.strip()}" for p in src.pages)
            (a.out / "text" / f"{name}.txt").write_text(body + "\n", encoding="utf-8")
        visual = [p for p in src.pages if p.reasons]
        sheets: list[Path] = []
        if src.image is not None:
            dst = a.out / "sheets" / f"{name}-1.jpg"
            shutil.copy(src.image, dst)
            sheets = [dst]
        elif visual and src.render_pdf:
            pick = [p.n for p in visual][: a.max_visual]
            if len(visual) > a.max_visual:
                src.note += f"（需要看圖的頁面有 {len(visual)} 頁，只渲染前 {a.max_visual} 頁；其餘請依需要用 --dpi 或手動處理）"
            sheets = make_sheets(name, render_pages(src.render_pdf, pick, work, a.dpi), a.out / "sheets")
        total_sheets += len(sheets)
        chars = sum(len(re.sub(r"\s", "", p.text)) for p in src.pages)
        vis = ", ".join(str(p.n) for p in visual) or "—"
        sh = ", ".join(f"`sheets/{s.name}`" for s in sheets) or "—"
        span = f"（p.{src.spec}）" if src.spec else ""
        lines.append(f"| {src.path.name}{span} | {src.kind} | {len(src.pages)} | {chars} | {vis} | {sh} |")
        if src.note:
            lines.append(f"| ↳ 注意 | {src.note} | | | | |")

    lines += ["", "## 需要看圖的原因", ""]
    for src in sources:
        for p in src.pages:
            if p.reasons:
                lines.append(f"- {src.path.name} p.{p.n}：{'；'.join(p.reasons)}")
    lines += ["", f"共 {total_sheets} 張 sheet。每張約 1,500–2,500 tokens；一頁一頁看則約每頁 1,000–1,500 tokens。"]
    (a.out / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    shutil.rmtree(work, ignore_errors=True)
    print((a.out / "report.md").read_text(encoding="utf-8"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
